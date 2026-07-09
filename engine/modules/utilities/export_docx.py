"""Export chat sessions and tactical facts to Word (.docx) with txt fallback."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from modules.gui.session_store import ChatSession
from modules.tactical_context import get_facts_store


def _try_docx():
    try:
        from docx import Document

        return Document
    except ImportError:
        return None


def export_session_docx(session: ChatSession, dest: Path) -> Path:
    """Write session transcript to .docx (or .txt if python-docx missing)."""
    dest = Path(dest)
    Document = _try_docx()
    if Document is None:
        dest = dest.with_suffix(".txt")
        lines = [
            f"Arden Chat — {session.title}",
            f"Exported {datetime.now(timezone.utc).isoformat()}",
            "",
        ]
        for msg in session.messages:
            lines.append(f"{msg.speaker}\n{msg.body}\n")
        dest.write_text("\n".join(lines), encoding="utf-8")
        return dest

    if dest.suffix.lower() != ".docx":
        dest = dest.with_suffix(".docx")
    doc = Document()
    doc.add_heading(f"Arden Chat — {session.title}", level=1)
    doc.add_paragraph(f"Exported {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    for msg in session.messages:
        doc.add_heading(msg.speaker, level=3)
        doc.add_paragraph(msg.body)
    doc.save(str(dest))
    return dest


def export_facts_docx(dest: Path) -> Path:
    """Export confirmed tactical facts."""
    facts = get_facts_store().list_facts(limit=500)
    dest = Path(dest)
    Document = _try_docx()
    if Document is None:
        dest = dest.with_suffix(".txt")
        lines = ["What I know about you", ""]
        for f in facts:
            domain = f.get("domain", "?").replace("_", " ").title()
            lines.append(f"[{domain}] {f.get('content', '')}")
        dest.write_text("\n".join(lines), encoding="utf-8")
        return dest

    if dest.suffix.lower() != ".docx":
        dest = dest.with_suffix(".docx")
    doc = Document()
    doc.add_heading("What I know about you", level=1)
    doc.add_paragraph("Confirmed facts stored on this PC.")
    current_domain = None
    for f in facts:
        domain = f.get("domain", "?").replace("_", " ").title()
        if domain != current_domain:
            doc.add_heading(domain, level=2)
            current_domain = domain
        doc.add_paragraph(f.get("content", ""), style="List Bullet")
    doc.save(str(dest))
    return dest
